"""Bloomberg parser must degrade gracefully on any workbook — never raise."""

import openpyxl

from app.fetchers import bloomberg_excel
from app.models.bloomberg import BloombergData


def _make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"] = "СЫРЬЕВОЙ РЫНОК"
    ws["A2"] = "name"; ws["B2"] = "price"
    ws["A3"] = "Золото"; ws["B3"] = 4119.59; ws["C3"] = -389.73
    ws["A4"] = "Нефть"; ws["B4"] = 77.26; ws["C4"] = -26.28
    wb.save(path)


def test_parses_anchored_block(tmp_path):
    p = tmp_path / "bb.xlsx"
    _make_workbook(p)
    data = bloomberg_excel.parse(str(p))
    assert isinstance(data, BloombergData)
    names = [c.name for c in data.commodities.rows]
    assert "Золото" in names and "Нефть" in names
    gold = next(c for c in data.commodities.rows if c.name == "Золото")
    assert gold.price == 4119.59
    assert gold.change_1m == -389.73


def test_missing_file_does_not_raise():
    data = bloomberg_excel.parse("/nonexistent/file.xlsx")
    assert isinstance(data, BloombergData)
    assert data.meta.status.value == "error"
    assert data.meta.warning
