"""Bloomberg .xlsx parser, driven entirely by config/bloomberg_map.yaml.

No cell positions are hardcoded here. Tables are located by an anchor label
and read via column offsets defined in the map. Anything missing is left
blank (still editable in the UI) rather than raising.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

import openpyxl
import yaml

from app.core.config import get_settings
from app.models.bloomberg import (
    BloombergData,
    Commodity,
    CommodityTable,
    EquityIndex,
    EquityIndexTable,
    RateCurveRow,
    RateCurveTable,
    TreasuryRow,
    TreasuryTable,
)
from app.models.common import FetchStatus, SourceMeta
from app.models.fx import FxCross, FxCrossTable


def _load_map(path: Optional[str] = None) -> dict:
    path = path or get_settings().bloomberg_map_path
    with open(path, "r", encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def _num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _find_anchor(ws, anchor: str) -> Optional[tuple[int, int]]:
    """Return (row, col) 1-based of the first cell containing `anchor`."""
    needle = anchor.casefold()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and needle in str(cell.value).casefold():
                return cell.row, cell.column
    return None


def _read_block(ws, spec: dict, defaults: dict) -> list[dict]:
    """Read rows under a spec's anchor into list of {field: rawvalue}."""
    anchor = _find_anchor(ws, spec["anchor"])
    if anchor is None:
        return []
    arow, acol = anchor
    start = arow + spec.get(
        "data_start_row_offset", defaults.get("data_start_row_offset", 2)
    )
    max_rows = spec.get("max_rows", defaults.get("max_rows", 30))
    columns: dict[str, int] = spec["columns"]

    out: list[dict] = []
    for r in range(start, start + max_rows):
        record: dict[str, Any] = {}
        empty = True
        for field, off in columns.items():
            if off is None:
                record[field] = None
                continue
            cell = ws.cell(row=r, column=acol + int(off))
            record[field] = cell.value
            if cell.value not in (None, ""):
                empty = False
        if empty:
            break  # blank row terminates the block
        out.append(record)
    return out


def parse(xlsx_path: str | Path, map_path: Optional[str] = None) -> BloombergData:
    fetched_at = None
    try:
        bmap = _load_map(map_path)
        defaults = bmap.get("defaults", {})
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

        def ws_for(spec: dict):
            name = spec["sheet"]
            return wb[name] if name in wb.sheetnames else wb[wb.sheetnames[0]]

        # Rate curves
        curves: list[RateCurveTable] = []
        for spec in bmap.get("rate_curves", []):
            recs = _read_block(ws_for(spec), spec, defaults)
            rows = [
                RateCurveRow(
                    date=_text(rec.get("date")),
                    on=_num(rec.get("on")),
                    m1=_num(rec.get("m1")),
                    m3=_num(rec.get("m3")),
                    m6=_num(rec.get("m6")),
                    m12=_num(rec.get("m12")),
                )
                for rec in recs
            ]
            curves.append(RateCurveTable(name=spec["name"], rows=rows))

        # FX crosses
        fx_spec = bmap.get("fx")
        fx_rows = []
        if fx_spec:
            for rec in _read_block(ws_for(fx_spec), fx_spec, defaults):
                fx_rows.append(
                    FxCross(
                        code=_text(rec.get("code")),
                        price=_num(rec.get("price")),
                        change_1m=_num(rec.get("change_1m")),
                        change_6m=_num(rec.get("change_6m")),
                        change_12m=_num(rec.get("change_12m")),
                    )
                )

        # Commodities
        commodities = []
        if bmap.get("commodities"):
            spec = bmap["commodities"]
            for rec in _read_block(ws_for(spec), spec, defaults):
                commodities.append(
                    Commodity(
                        name=_text(rec.get("name")),
                        price=_num(rec.get("price")),
                        change_1m=_num(rec.get("change_1m")),
                        change_6m=_num(rec.get("change_6m")),
                        change_12m=_num(rec.get("change_12m")),
                    )
                )

        # Equities
        equities = []
        if bmap.get("equities"):
            spec = bmap["equities"]
            for rec in _read_block(ws_for(spec), spec, defaults):
                equities.append(
                    EquityIndex(
                        name=_text(rec.get("name")),
                        price=_num(rec.get("price")),
                        change_1m=_num(rec.get("change_1m")),
                        change_6m=_num(rec.get("change_6m")),
                        change_12m=_num(rec.get("change_12m")),
                    )
                )

        # Treasuries
        treasuries = []
        if bmap.get("treasuries"):
            spec = bmap["treasuries"]
            for rec in _read_block(ws_for(spec), spec, defaults):
                treasuries.append(
                    TreasuryRow(
                        tenor=_text(rec.get("tenor")),
                        yld=_num(rec.get("yld")),
                        change_1m=_num(rec.get("change_1m")),
                        change_6m=_num(rec.get("change_6m")),
                        change_12m=_num(rec.get("change_12m")),
                    )
                )
        wb.close()

        found_any = any([curves, fx_rows, commodities, equities, treasuries])
        meta = SourceMeta(
            source=f"Bloomberg export: {Path(xlsx_path).name}",
            status=FetchStatus.OK if found_any else FetchStatus.EMPTY,
            warning=None
            if found_any
            else "No tables matched bloomberg_map.yaml anchors — check the map.",
            fetched_at=fetched_at,
        )
        return BloombergData(
            rate_curves=curves,
            fx=FxCrossTable(rows=fx_rows, meta=meta),
            commodities=CommodityTable(rows=commodities),
            equities=EquityIndexTable(rows=equities),
            treasuries=TreasuryTable(rows=treasuries),
            meta=meta,
        )

    except Exception as exc:  # noqa: BLE001 - never crash generation
        meta = SourceMeta(
            source="Bloomberg export",
            status=FetchStatus.ERROR,
            warning=f"Failed to parse workbook: {exc}",
        )
        return BloombergData(
            fx=FxCrossTable(rows=[], meta=meta),
            commodities=CommodityTable(),
            equities=EquityIndexTable(),
            treasuries=TreasuryTable(),
            meta=meta,
        )
